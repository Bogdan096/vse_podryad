import requests
from bs4 import BeautifulSoup
import asyncio
import itertools
import logging
import pandas as pd
from aiogram.handlers import message
from entrance_data import api_id, api_hash
from telethon import TelegramClient, events
import os
import openpyxl



'''брать из файла excel пользователей (2 лист), отправка текстом в чат, gitignore, решение двухфакторного входа, отправка по времени


session = "test"
api_id =  24518272
api_hash = '7d643a28aaffaab9a34a04f36da6d44c'''

session = "test"
unames = []
spisok = []
alive_bots = []
id_list = []
cl = []
pop_cl = []
slovar = {"Боты":cl}
client = TelegramClient(session,api_id, api_hash,system_version="4.16.30-vxCUSTOM")
getter_name = []
chats = tuple(id_list)
flag = 0


@client.on(events.NewMessage())
async def handle_message(event):
        print("message handled")
        sender_username =  await event.get_sender()
        cl.remove(sender_username.first_name)
        for uname in getter_name:
            username = "Start_Checker_bot"
            entit = await client.get_entity(username)
            await client.send_message(entity=entit, message=f"{sender_username.first_name} активен")



def name_parse():
    print("it works")
    path = os.path.abspath("../Боты и скрипты.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb.sheetnames
    urls = ws[0]
    getters = ws[1]
    print("it workz")
    for i in range(0, wb[urls].max_row):
        for col in wb[urls].iter_cols(1, 3):
            if "https" in str(col[i].value):
                spisok.append(str(col[i].value))
    for i in range(1, wb[getters].max_row):
        for col in wb[getters].iter_cols(2):
            getter_name.append(str(col[i].value))
    for link in spisok:
        resp = requests.get(link)
        soup = BeautifulSoup(resp.text, "lxml")
        name = soup.find(class_="tgme_page_title")
        username = soup.find(class_="tgme_page_extra")
        cl.append(name.text)
        print(cl)
        id_list.append(username.text[4:])

# async def make_file(spisok: list):
#     path = os.path.abspath("../Отчет о ботах.xlsx")
#     if os.path.exists(path):
#         os.remove(path)
#     alive_bots.clear()
#     for response_uname in chats:
#         if response_uname in id_list:
#             alive_bots.append("Работает")
#         else:
#             alive_bots.append("Не работает")
#     slovar["Статус"] = alive_bots
#     # print(slovar)
#     df = pd.DataFrame(slovar)
#     print(df)
#     df.to_excel("Отчет о ботах.xlsx", index=False)
#     username = ""
#     entit = await client.get_entity(username)
#     await client.send_message(entity=entit, file="Отчет о ботах.xlsx")


async def handle_messages_to_send():

    for elem in ["Start_Checker_bot","Bodyapkin_bot"]:
        username = elem
        entit = await client.get_entity(username)
        await client.send_message(entity=entit, message= "/start")

async def main():
    name_parse()
    await handle_messages_to_send()

    await client.run_until_disconnected()


if __name__ == '__main__':
    with client:

        client.loop.run_until_complete(
            main()
        )
